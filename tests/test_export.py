"""Tests for export.py — verifies Excel structure without network calls."""
import os
import tempfile

import pytest
from openpyxl import load_workbook

from export import _bytes_to_human, write_traffic_report

SAMPLE_RESULTS = {
    "By Hostname": [
        {
            "hostname": "www.example.com",
            "edgeBytesSum": 10_737_418_240,  # 10 GB
            "edgeHitsSum": 500000,
            "originBytesSum": 1_073_741_824,  # 1 GB
            "originHitsSum": 50000,
            "midgressBytesSum": 0,
        }
    ],
    "By CP Code": [
        {
            "cpcode": 12345,
            "edgeBytesSum": 5_368_709_120,
            "originBytesSum": 536_870_912,
            "midgressBytesSum": 0,
            "offloadedBytesPercentage": 0.90,
        }
    ],
    "HTTP Status": [
        {"responseClass": "4xx", "edgeHitsSum": 1200, "originHitsSum": 300},
        {"responseClass": "5xx", "edgeHitsSum": 50, "originHitsSum": 10},
    ],
    "Offload": [
        {
            "cpcode": 12345,
            "offloadedBytesPercentage": 0.90,
            "offloadedHitsPercentage": 0.92,
            "edgeBytesSum": 5_368_709_120,
            "originBytesSum": 536_870_912,
        }
    ],
}


@pytest.fixture
def report_path(tmp_path):
    return str(tmp_path / "test_report.xlsx")


def test_write_creates_file(report_path):
    path = write_traffic_report(
        SAMPLE_RESULTS,
        report_path,
        account_name="Test Account",
        start="2025-01-01T00:00:00Z",
        end="2025-01-15T00:00:00Z",
    )
    assert os.path.exists(path)


def test_correct_sheet_names(report_path):
    write_traffic_report(
        SAMPLE_RESULTS,
        report_path,
        account_name="Test Account",
        start="2025-01-01T00:00:00Z",
        end="2025-01-15T00:00:00Z",
    )
    wb = load_workbook(report_path)
    assert "Summary" in wb.sheetnames
    assert "By Hostname" in wb.sheetnames
    assert "By CP Code" in wb.sheetnames
    assert "HTTP Status" in wb.sheetnames
    assert "Offload" in wb.sheetnames


def test_cover_sheet_has_account_name(report_path):
    write_traffic_report(
        SAMPLE_RESULTS,
        report_path,
        account_name="Acme Corp",
        start="2025-01-01T00:00:00Z",
        end="2025-01-15T00:00:00Z",
    )
    wb = load_workbook(report_path)
    ws = wb["Summary"]
    all_values = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("Acme Corp" in v for v in all_values)


def test_data_sheet_has_header_row(report_path):
    write_traffic_report(
        SAMPLE_RESULTS,
        report_path,
        account_name="Test",
        start="2025-01-01T00:00:00Z",
        end="2025-01-15T00:00:00Z",
    )
    wb = load_workbook(report_path)
    ws = wb["By Hostname"]
    # Header row is row 3; check column names are present
    header_values = [c.value for c in ws[3] if c.value]
    assert "hostname" in header_values
    assert "edgeBytesSum" in header_values


def test_returns_absolute_path(report_path):
    result = write_traffic_report(
        SAMPLE_RESULTS,
        report_path,
        account_name="Test",
        start="2025-01-01T00:00:00Z",
        end="2025-01-15T00:00:00Z",
    )
    assert os.path.isabs(result)


def test_empty_data_sheet_handled(tmp_path):
    path = str(tmp_path / "empty.xlsx")
    write_traffic_report(
        {"By Hostname": [], "By CP Code": []},
        path,
        account_name="Test",
        start="2025-01-01T00:00:00Z",
        end="2025-01-15T00:00:00Z",
    )
    wb = load_workbook(path)
    ws = wb["By Hostname"]
    assert ws["A1"].value == "No data returned for this report."


# ── _bytes_to_human helper ────────────────────────────────────

def test_bytes_to_human_gb():
    assert "GB" in _bytes_to_human(10_737_418_240.0)


def test_bytes_to_human_tb():
    assert "TB" in _bytes_to_human(1_099_511_627_776.0)


def test_bytes_to_human_mb():
    assert "MB" in _bytes_to_human(5_242_880.0)
